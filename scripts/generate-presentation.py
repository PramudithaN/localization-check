import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 Widescreen dimensions
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Professional Color Palette (Clean Light Theme, Solid Colors Only, Zero Gradients)
BG_WHITE = RGBColor(255, 255, 255)
BG_LIGHT_SLATE = RGBColor(248, 250, 252)       # #F8FAFC
BG_CARD_BLUE = RGBColor(241, 245, 249)         # #F1F5F9
BG_SOFT_RED = RGBColor(254, 242, 242)          # #FEF2F2
BG_SOFT_GREEN = RGBColor(240, 253, 244)        # #F0FDF4
BG_CODE_BOX = RGBColor(241, 245, 249)

BORDER_LIGHT = RGBColor(226, 232, 240)         # #E2E8F0
BORDER_BLUE = RGBColor(191, 219, 254)          # #BFDBFE
BORDER_GREEN = RGBColor(187, 247, 208)         # #BBF7D0
BORDER_RED = RGBColor(254, 202, 202)           # #FECACA

TEXT_NAVY = RGBColor(15, 23, 42)               # #0F172A (Primary Heading)
TEXT_BODY = RGBColor(51, 65, 85)               # #334155 (Primary Body)
TEXT_MUTED = RGBColor(100, 116, 139)           # #64748B (Secondary / Meta)
TEXT_BLUE = RGBColor(37, 99, 235)              # #2563EB (Accent Blue)
TEXT_GREEN = RGBColor(5, 150, 105)             # #059669 (Accent Green)
TEXT_RED = RGBColor(220, 38, 38)               # #DC2626 (Accent Red)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_CODE = "Consolas"

def set_slide_background(slide, color):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = color
    bg_shape.line.color.rgb = color
    bg_shape.line.width = Pt(0)
    return bg_shape

def add_header(slide, category_text, title_text, subtitle_text=None):
    # Category Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_text.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = TEXT_BLUE

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_NAVY

    # Subtitle (optional)
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=BG_LIGHT_SLATE, border_color=BORDER_LIGHT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

# ==============================================================================
# SLIDE 1: Title Slide (Human, Clean, Confident)
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1, BG_WHITE)

# Left hero box
hero_bg = add_card(slide1, Inches(0.8), Inches(0.8), Inches(7.5), Inches(5.9), BG_LIGHT_SLATE, BORDER_LIGHT)

hero_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(6.7), Inches(5.1))
tf1 = hero_box.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p_badge = tf1.paragraphs[0]
p_badge.text = "VS CODE EXTENSION | OPEN SOURCE (MIT)"
p_badge.font.name = FONT_HEADING
p_badge.font.size = Pt(11)
p_badge.font.bold = True
p_badge.font.color.rgb = TEXT_BLUE
p_badge.space_after = Pt(18)

p_h1 = tf1.add_paragraph()
p_h1.text = "Localization Check"
p_h1.font.name = FONT_HEADING
p_h1.font.size = Pt(36)
p_h1.font.bold = True
p_h1.font.color.rgb = TEXT_NAVY
p_h1.space_after = Pt(10)

p_h2 = tf1.add_paragraph()
p_h2.text = "Catching hardcoded UI text before our users do."
p_h2.font.name = FONT_HEADING
p_h2.font.size = Pt(18)
p_h2.font.color.rgb = TEXT_BODY
p_h2.space_after = Pt(18)

p_desc = tf1.add_paragraph()
p_desc.text = (
    "A developer tool that flags unlocalized strings in changed TypeScript and React files. "
    "It uses AST compiler trees instead of noisy regex, and fixes them with Copilot in one keystroke."
)
p_desc.font.name = FONT_BODY
p_desc.font.size = Pt(13)
p_desc.font.color.rgb = TEXT_MUTED
p_desc.space_after = Pt(28)

p_meta = tf1.add_paragraph()
p_meta.text = "Works with: React | Next.js | React Native | TypeScript | JavaScript"
p_meta.font.name = FONT_CODE
p_meta.font.size = Pt(10.5)
p_meta.font.color.rgb = TEXT_BODY

# Right side 3 stat/highlight cards
stat_items = [
    ("No RegEx Guesswork", "Uses the Babel AST parser to inspect actual syntax nodes. It knows the difference between a button label and a CSS class name.", BG_WHITE, BORDER_LIGHT),
    ("Copilot Automated Fix", "Press Alt + L on any underlined string. Copilot inserts the hook, updates en.json, and writes clean code for you in 2 seconds.", BG_WHITE, BORDER_LIGHT),
    ("Zero Noise for Clean Files", "Only runs live checks on files modified in Git or unsaved buffers. It never bogs down huge enterprise repositories.", BG_WHITE, BORDER_LIGHT)
]

for idx, (head, desc, bg, bdr) in enumerate(stat_items):
    c_top = Inches(0.8 + idx * 2.05)
    add_card(slide1, Inches(8.55), c_top, Inches(4.0), Inches(1.8), bg, bdr)
    
    tb = slide1.shapes.add_textbox(Inches(8.8), c_top + Inches(0.2), Inches(3.5), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = head
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_NAVY
    p1.space_after = Pt(6)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 2: The Real Problem (Human, Honest, Relatable)
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2, BG_WHITE)
add_header(slide2, "The Real Problem", "Why hardcoded strings sneak into production", "Every frontend team deals with this, especially when shipping under deadlines.")

col_w = Inches(3.75)
gap = Inches(0.25)

cards_data = [
    (
        "1. The Sprint Crunch",
        "A developer is shipping a modal at 6 PM on Thursday. They write <button>Save Details</button> or toast.error('Payment failed'), intending to localize it tomorrow. It gets approved in PR, and next week our Spanish users see English buttons.",
        BG_SOFT_RED, BORDER_RED, TEXT_RED,
        "Human factor: Nobody forgets on purpose; manual steps get skipped when rushing."
    ),
    (
        "2. Noisy Linters Cry Wolf",
        "Standard regex linters scream at everything: className='text-red-500', testId='user-id', or color='blue'. Developers get notification fatigue, ignore the warnings, or just disable the linter rule altogether.",
        BG_CARD_BLUE, BORDER_BLUE, TEXT_BLUE,
        "Linter trap: When a tool gives 50 false alarms, developers stop trusting it."
    ),
    (
        "3. The en.json Tax",
        "Even when doing it right, the workflow is annoying: think of a key name, switch to locales/en.json, paste the key, switch back, write the hook, import react-i18next. Doing this 40 times in a feature sprint drains mental energy.",
        BG_SOFT_GREEN, BORDER_GREEN, TEXT_GREEN,
        "Friction tax: High-friction processes invite shortcuts and missed keys."
    )
]

for i, (title, body, bg, bdr, accent_col, footer_note) in enumerate(cards_data):
    left_pos = Inches(0.8 + i * (3.75 + 0.22))
    add_card(slide2, left_pos, Inches(1.85), col_w, Inches(4.85), bg, bdr)
    
    tb = slide2.shapes.add_textbox(left_pos + Inches(0.3), Inches(2.15), col_w - Inches(0.6), Inches(4.25))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(17)
    p_t.font.bold = True
    p_t.font.color.rgb = accent_col
    p_t.space_after = Pt(14)
    
    p_b = tf.add_paragraph()
    p_b.text = body
    p_b.font.name = FONT_BODY
    p_b.font.size = Pt(12)
    p_b.font.color.rgb = TEXT_BODY
    p_b.space_after = Pt(20)
    
    p_f = tf.add_paragraph()
    p_f.text = footer_note
    p_f.font.name = FONT_BODY
    p_f.font.size = Pt(11)
    p_f.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 3: How the Engine Works (Precision & Clarity)
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3, BG_WHITE)
add_header(slide3, "Under The Hood", "Real AST parsing, not pattern matching", "The extension parses your TypeScript code the exact same way a compiler does.")

# Left Column: What It Checks
add_card(slide3, Inches(0.8), Inches(1.85), Inches(5.75), Inches(4.85), BG_LIGHT_SLATE, BORDER_LIGHT)
tb_left = slide3.shapes.add_textbox(Inches(1.15), Inches(2.15), Inches(5.05), Inches(4.25))
tf_l = tb_left.text_frame
tf_l.word_wrap = True
tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0

p_lh = tf_l.paragraphs[0]
p_lh.text = "What It Flags (High Confidence)"
p_lh.font.name = FONT_HEADING
p_lh.font.size = Pt(16)
p_lh.font.bold = True
p_lh.font.color.rgb = TEXT_NAVY
p_lh.space_after = Pt(12)

checks = [
    ("JSX Text & Tags:", "Headings, paragraphs, spans, button titles (e.g. <button>Submit</button>)"),
    ("UI Attributes:", "label, title, placeholder, tooltip, aria-label, alt, helperText"),
    ("Object Properties:", "title: 'Edit Profile', label: 'Email', message: 'Saved successfully'"),
    ("Notification Calls:", "showNotification('warning', 'Session expired') — skips the status type!"),
    ("Template Literals:", "Flags `Welcome back, ${user.name}!` while ignoring technical routes")
]

for title, desc in checks:
    p = tf_l.add_paragraph()
    p.text = f"• {title} "
    p.font.name = FONT_HEADING
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_NAVY
    
    run = p.add_run()
    run.text = desc
    run.font.name = FONT_BODY
    run.font.bold = False
    run.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

# Right Column: What It Ignores
add_card(slide3, Inches(6.8), Inches(1.85), Inches(5.75), Inches(4.85), BG_LIGHT_SLATE, BORDER_LIGHT)
tb_right = slide3.shapes.add_textbox(Inches(7.15), Inches(2.15), Inches(5.05), Inches(4.25))
tf_r = tb_right.text_frame
tf_r.word_wrap = True
tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0

p_rh = tf_r.paragraphs[0]
p_rh.text = "What It Safely Ignores (Zero Noise)"
p_rh.font.name = FONT_HEADING
p_rh.font.size = Pt(16)
p_rh.font.bold = True
p_rh.font.color.rgb = TEXT_NAVY
p_rh.space_after = Pt(12)

ignores = [
    ("50+ Programming Types:", "Promise, Observable, HTMLElement, Record, AxiosResponse, etc."),
    ("Styling & Props:", "className, id, style, data-testid, color, width, align, variant"),
    ("Routes & URLs:", "https://api.domain.com, /dashboard/${tenantId}/settings"),
    ("Technical Blocks:", "spec: { label: 'batchTypeName', value: 'id' } in data mappings"),
    ("Ignored Code Elements:", "Any code or text inside <code>, <pre>, <script>, and <style>")
]

for title, desc in ignores:
    p = tf_r.add_paragraph()
    p.text = f"✓ {title} "
    p.font.name = FONT_HEADING
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_GREEN
    
    run = p.add_run()
    run.text = desc
    run.font.name = FONT_BODY
    run.font.bold = False
    run.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

# Bottom note inside right card
p_foot = tf_r.add_paragraph()
p_foot.text = "Net result: When a squiggly line shows up, it's actually something a human needs to fix."
p_foot.font.name = FONT_BODY
p_foot.font.size = Pt(11)
p_foot.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 4: The Copilot Integration (Concrete Before/After)
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4, BG_WHITE)
add_header(slide4, "Copilot AI Workflow", "Press Alt + L: Done in 2 seconds", "Turn tedious copy-pasting into a one-keystroke action.")

# Before Box
add_card(slide4, Inches(0.8), Inches(1.85), Inches(5.75), Inches(4.85), BG_CARD_BLUE, BORDER_LIGHT)
tb_b4 = slide4.shapes.add_textbox(Inches(1.15), Inches(2.15), Inches(5.05), Inches(4.25))
tf_b4 = tb_b4.text_frame
tf_b4.word_wrap = True
tf_b4.margin_left = tf_b4.margin_top = tf_b4.margin_right = tf_b4.margin_bottom = 0

p_b4_h = tf_b4.paragraphs[0]
p_b4_h.text = "BEFORE (What the developer wrote)"
p_b4_h.font.name = FONT_HEADING
p_b4_h.font.size = Pt(14)
p_b4_h.font.bold = True
p_b4_h.font.color.rgb = TEXT_RED
p_b4_h.space_after = Pt(12)

p_code_b4 = tf_b4.add_paragraph()
p_code_b4.text = (
    "export function DeleteModal() {\n"
    "    return (\n"
    "        <div>\n"
    "            <h3>Delete Account</h3>\n"
    "            <button title=\"Cancel modal\">\n"
    "                Discard changes\n"
    "            </button>\n"
    "        </div>\n"
    "    );\n"
    "}"
)
p_code_b4.font.name = FONT_CODE
p_code_b4.font.size = Pt(11)
p_code_b4.font.color.rgb = TEXT_NAVY
p_code_b4.space_after = Pt(14)

p_b4_note = tf_b4.add_paragraph()
p_b4_note.text = "Notice: 'Delete Account', 'Cancel modal', and 'Discard changes' are all hardcoded in English."
p_b4_note.font.name = FONT_BODY
p_b4_note.font.size = Pt(11)
p_b4_note.font.color.rgb = TEXT_MUTED

# After Box
add_card(slide4, Inches(6.8), Inches(1.85), Inches(5.75), Inches(4.85), BG_SOFT_GREEN, BORDER_GREEN)
tb_aft = slide4.shapes.add_textbox(Inches(7.15), Inches(2.15), Inches(5.05), Inches(4.25))
tf_aft = tb_aft.text_frame
tf_aft.word_wrap = True
tf_aft.margin_left = tf_aft.margin_top = tf_aft.margin_right = tf_aft.margin_bottom = 0

p_aft_h = tf_aft.paragraphs[0]
p_aft_h.text = "AFTER (Alt + L or Alt + Shift + L)"
p_aft_h.font.name = FONT_HEADING
p_aft_h.font.size = Pt(14)
p_aft_h.font.bold = True
p_aft_h.font.color.rgb = TEXT_GREEN
p_aft_h.space_after = Pt(12)

p_code_aft = tf_aft.add_paragraph()
p_code_aft.text = (
    "import { useTranslation } from 'react-i18next'; // Auto-inserted\n\n"
    "export function DeleteModal() {\n"
    "    const { t } = useTranslation();              // Auto-injected\n"
    "    return (\n"
    "        <div>\n"
    "            <h3>{t('common.deleteAccount')}</h3>\n"
    "            <button title={t('common.cancelModal')}>\n"
    "                {t('common.discard')}\n"
    "            </button>\n"
    "        </div>\n"
    "    );\n"
    "}"
)
p_code_aft.font.name = FONT_CODE
p_code_aft.font.size = Pt(10)
p_code_aft.font.color.rgb = TEXT_NAVY
p_code_aft.space_after = Pt(12)

p_aft_note = tf_aft.add_paragraph()
p_aft_note.text = "And en.json was automatically updated with 'common.deleteAccount', 'common.cancelModal', and 'common.discard' without you ever opening the file."
p_aft_note.font.name = FONT_BODY
p_aft_note.font.size = Pt(11)
p_aft_note.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 5: Thoughtful Engineering (Built to Stay Out of Your Way)
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5, BG_WHITE)
add_header(slide5, "Developer Experience", "Designed to be helpful, not in your face", "Small details that make this tool a joy to use on real teams.")

features = [
    (
        "Git-Staged Only",
        "It doesn't scan your 2,000 legacy files every time you open VS Code. It only scans lines you touched in your Git branch or unsaved buffers. Zero CPU spin.",
        "Alt + S to force re-scan"
    ),
    (
        "Pre-Commit Heads-Up",
        "If you stage a commit with unlocalized strings, a discreet VS Code toast pops up. You can inspect it with one click before locking in the git commit.",
        "Alt + R for full staged check"
    ),
    (
        "Dead Key Finder",
        "Over time, old features get deleted but their keys stay in en.json forever. Press Alt + U and the tool statically scans your whole codebase to find unused keys.",
        "Alt + U audits en.json"
    ),
    (
        "1-Click Learning",
        "Have a custom component like <AlertBanner text='...' />? Hit Alt + F at your cursor to teach the extension to check that prop across your team forever.",
        "Alt + F to flag / Alt + M to ignore"
    )
]

for idx, (title, body, shortcut) in enumerate(features):
    c_left = Inches(0.8 + idx * (2.85 + 0.13))
    add_card(slide5, c_left, Inches(1.85), Inches(2.85), Inches(4.85), BG_LIGHT_SLATE, BORDER_LIGHT)
    
    tb = slide5.shapes.add_textbox(c_left + Inches(0.25), Inches(2.15), Inches(2.35), Inches(4.25))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(16)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_NAVY
    p_t.space_after = Pt(12)
    
    p_b = tf.add_paragraph()
    p_b.text = body
    p_b.font.name = FONT_BODY
    p_b.font.size = Pt(11.5)
    p_b.font.color.rgb = TEXT_BODY
    p_b.space_after = Pt(24)
    
    p_s = tf.add_paragraph()
    p_s.text = shortcut
    p_s.font.name = FONT_CODE
    p_s.font.size = Pt(10)
    p_s.font.bold = True
    p_s.font.color.rgb = TEXT_BLUE

# ==============================================================================
# SLIDE 6: Summary & Practical Takeaway
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6, BG_WHITE)
add_header(slide6, "The Bottom Line", "Why engineering teams actually keep this installed", "Better international user experience without slowing down feature delivery.")

add_card(slide6, Inches(0.8), Inches(1.85), Inches(11.733), Inches(4.85), BG_LIGHT_SLATE, BORDER_LIGHT)
tb_s6 = slide6.shapes.add_textbox(Inches(1.3), Inches(2.25), Inches(10.7), Inches(4.0))
tf_s6 = tb_s6.text_frame
tf_s6.word_wrap = True
tf_s6.margin_left = tf_s6.margin_top = tf_s6.margin_right = tf_s6.margin_bottom = 0

p_s6_lead = tf_s6.paragraphs[0]
p_s6_lead.text = "In short:"
p_s6_lead.font.name = FONT_HEADING
p_s6_lead.font.size = Pt(20)
p_s6_lead.font.bold = True
p_s6_lead.font.color.rgb = TEXT_NAVY
p_s6_lead.space_after = Pt(16)

points = [
    ("No more embarrassing English strings in localized markets.", "Catches errors right where they happen: in the code editor, while you're writing."),
    ("90% less manual grunt work.", "Alt + L automates the entire i18n plumbing (imports, hooks, keys, and dictionary updates)."),
    ("Zero runtime overhead in production.", "It is strictly a developer-time tool (VS Code extension). No npm runtime dependencies in your bundle."),
    ("Built to be trusted.", "21 automated test cases, zero false alarms on standard code patterns, and safe reverse-order editing.")
]

for lead, sub in points:
    p = tf_s6.add_paragraph()
    p.text = f"• {lead} "
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_NAVY
    
    r = p.add_run()
    r.text = sub
    r.font.name = FONT_BODY
    r.font.bold = False
    r.font.color.rgb = TEXT_BODY
    p.space_after = Pt(12)

p_closing = tf_s6.add_paragraph()
p_closing.text = "Ready to install: Available on VS Code Marketplace & VSIX package (v0.5.9)"
p_closing.font.name = FONT_CODE
p_closing.font.size = Pt(11.5)
p_closing.font.bold = True
p_closing.font.color.rgb = TEXT_BLUE
p_closing.space_before = Pt(8)

# Save presentation
output_path = "Localization_Check_Presentation.pptx"
prs.save(output_path)
print(f"Successfully generated PowerPoint presentation: {output_path}")
